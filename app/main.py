import asyncio
import io
import json
import logging
import os
from pathlib import Path
import secrets
import sqlite3
import time

from fastapi import FastAPI, UploadFile, HTTPException, Request
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from docx import Document
import pymupdf
from app.llm import KINDS, ROOT, model_name, ExtractionError, provider_name, validate_configuration
from app.persons import validate_spans, person_label
from app.extract_persons import extract as persons
from app.extract_organizations import extract as organizations
from app.extract_entrepreneurs import extract as entrepreneurs
from app.extract_communications import extract as communications

app = FastAPI(title="Проверка сущностей")
app.mount("/static", StaticFiles(directory=ROOT / "app/static"), name="static")
DATA = Path(os.getenv("DATA_DIR", str(ROOT / "data")))
DATA.mkdir(parents=True, exist_ok=True)
JOBS = {}
TTL = 3600

def database():
    db = sqlite3.connect(DATA / "entities.sqlite3")
    db.execute("PRAGMA foreign_keys=ON")
    db.executescript("""
    CREATE TABLE IF NOT EXISTS documents(id INTEGER PRIMARY KEY, review_token TEXT UNIQUE NOT NULL,
    filename TEXT NOT NULL, format TEXT NOT NULL, text TEXT NOT NULL, uploaded_at REAL NOT NULL,
    confirmed_at REAL NOT NULL, operator TEXT, model TEXT NOT NULL, original_result TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS entities(id INTEGER PRIMARY KEY, document_id INTEGER NOT NULL REFERENCES documents(id),
    kind TEXT NOT NULL, normalized_value TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS mentions(id INTEGER PRIMARY KEY, entity_id INTEGER NOT NULL REFERENCES entities(id),
    original_text TEXT NOT NULL, start INTEGER NOT NULL, end INTEGER NOT NULL, note TEXT);
    """)
    # Additive migration: keep earlier documents and mentions intact.
    if "field" not in {row[1] for row in db.execute("PRAGMA table_info(mentions)")}:
        db.execute("ALTER TABLE mentions ADD COLUMN field TEXT NOT NULL DEFAULT 'entity'")
    db.execute("""CREATE TABLE IF NOT EXISTS person_details(
        entity_id INTEGER PRIMARY KEY REFERENCES entities(id), surname TEXT, given_name TEXT,
        patronymic TEXT, birth_date TEXT, birth_place TEXT)""")
    if "provider" not in {row[1] for row in db.execute("PRAGMA table_info(documents)")}:
        db.execute("ALTER TABLE documents ADD COLUMN provider TEXT NOT NULL DEFAULT 'local'")
    db.commit()
    return db

def prune():
    for token, job in list(JOBS.items()):
        if time.time() - job["uploaded_at"] > TTL and job["status"] != "processing":
            del JOBS[token]

def read_text(filename, contents):
    suffix = Path(filename).suffix.lower()
    if suffix == ".txt":
        try:
            return contents.decode("utf-8-sig")
        except UnicodeDecodeError:
            return contents.decode("cp1251")
    if suffix == ".docx":
        doc = Document(io.BytesIO(contents))
        paragraphs = [p.text for p in doc.paragraphs]
        paragraphs.extend(" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows)
        return "\n".join(paragraphs)
    if suffix == ".pdf":
        with pymupdf.open(stream=contents, filetype="pdf") as doc:
            return "\n".join(page.get_text() for page in doc)
    raise ValueError("Поддерживаются TXT, DOCX и PDF с текстовым слоем")

async def process(token):
    job = JOBS[token]
    try:
        found = []
        for stage, extractor in (("Сведения о людях", persons), ("Организации", organizations),
                                 ("ИП", entrepreneurs), ("Коммуникации", communications)):
            job["stage"] = stage
            found.extend(await extractor(job["text"]))
        unique = {(e["kind"], e["start"], e["end"], e["normalized_value"],
                   json.dumps(e.get("person"), sort_keys=True, ensure_ascii=False),
                   json.dumps(e.get("source_spans"), sort_keys=True)): e for e in found}
        job["entities"] = sorted(unique.values(), key=lambda e: (e["start"], e["kind"]))
        job["status"] = "review"
    except ExtractionError as exc:
        job["status"] = "error"
        job["error"] = f"{job.get('stage', 'Распознавание')}: {exc}"
    except Exception as exc:
        job["status"] = "error"
        job["error"] = "Внутренняя ошибка обработки. Тип ошибки указан в консоли Python сервера."
        logging.getLogger(__name__).error("Entity processing failed (%s)", type(exc).__name__)

@app.get("/")
def index():
    return FileResponse(ROOT / "app/templates/index.html")

@app.get("/api/config")
def configuration():
    provider = provider_name()
    return {"provider": provider,
            "notice": "Gemini API: текст документа отправляется Google для распознавания." if provider == "gemini"
                      else "Локальная модель: текст передаётся настроенному LLM серверу."}

@app.post("/api/upload")
async def upload(file: UploadFile):
    prune()
    if len(JOBS) >= 20:
        raise HTTPException(429, "Слишком много временных документов")
    contents = await file.read(20 * 1024 * 1024 + 1)
    if len(contents) > 20 * 1024 * 1024:
        raise HTTPException(413, "Максимальный размер файла — 20 МБ")
    filename = Path(file.filename or "document").name
    try:
        text = await asyncio.to_thread(read_text, filename, contents)
    except Exception as exc:
        raise HTTPException(400, "Не удалось прочитать файл. Поддерживаются TXT, DOCX и текстовый PDF") from exc
    if not text.strip():
        raise HTTPException(400, "Из файла не извлечён текст. Проверьте его содержимое; PDF-скану требуется OCR.")
    if len(text) > 300000:
        raise HTTPException(400, f"Извлечено {len(text)} символов; максимум — 300 000.")
    try:
        model = validate_configuration()
    except ValueError as exc:
        raise HTTPException(503, str(exc)) from exc
    token = secrets.token_urlsafe(32)
    JOBS[token] = {"filename": filename, "text": text, "uploaded_at": time.time(), "status": "processing",
                   "entities": [], "model": model, "provider": provider_name()}
    asyncio.create_task(process(token))
    return {"token": token}

@app.get("/api/review/{token}")
def review(token: str):
    prune()
    if token not in JOBS:
        raise HTTPException(404, "Временный документ не найден или срок проверки истёк")
    return JOBS[token]

@app.delete("/api/review/{token}")
def cancel(token: str):
    job = JOBS.get(token)
    if job and job["status"] == "processing":
        raise HTTPException(409, "Дождитесь завершения обработки")
    JOBS.pop(token, None)
    return {"cancelled": True}

@app.post("/api/review/{token}/confirm")
async def confirm(token: str, request: Request):
    prune()
    job = JOBS.get(token)
    if not job or job["status"] != "review":
        raise HTTPException(409, "Документ не готов к подтверждению")
    try:
        payload = await request.json()
    except ValueError as exc:
        raise HTTPException(400, "Неверный JSON") from exc
    entities = payload.get("entities") if isinstance(payload, dict) else None
    if not isinstance(entities, list) or len(entities) > 10000:
        raise HTTPException(400, "Неверный список сущностей")
    for e in entities:
        if not isinstance(e, dict):
            raise HTTPException(400, "Неверная сущность")
        start, end = e.get("start"), e.get("end")
        if (type(start) is not int or type(end) is not int or not 0 <= start < end <= len(job["text"])
            or e.get("kind") not in KINDS or not isinstance(e.get("normalized_value"), str)
            or not e["normalized_value"].strip()):
            raise HTTPException(400, "Проверьте категорию, границы и нормализованное значение")
        if e["kind"] == "persons":
            try:
                validate_spans(e.get("person"), e.get("source_spans"), job["text"])
            except ValueError as exc:
                raise HTTPException(400, str(exc)) from exc
            anchor = next(span for span in e["source_spans"].values() if span)
            e["start"], e["end"] = anchor["start"], anchor["end"]
            e["normalized_value"] = person_label(e["person"], job["text"][e["start"]:e["end"]])
    with database() as db:
        existing = db.execute("SELECT id FROM documents WHERE review_token=?", (token,)).fetchone()
        if existing:
            return {"document_id": existing[0], "already_saved": True}
        cursor = db.execute("INSERT INTO documents(review_token,filename,format,text,uploaded_at,confirmed_at,operator,model,original_result,provider) VALUES(?,?,?,?,?,?,?,?,?,?)",
            (token, job["filename"], Path(job["filename"]).suffix, job["text"], job["uploaded_at"], time.time(), None, job["model"], json.dumps(job["entities"], ensure_ascii=False), job.get("provider", "local")))
        document_id = cursor.lastrowid
        for e in entities:
            entity_id = db.execute("INSERT INTO entities(document_id,kind,normalized_value) VALUES(?,?,?)",
                (document_id, e["kind"], e["normalized_value"].strip())).lastrowid
            spans = {"entity": {"start": e["start"], "end": e["end"]}}
            if e["kind"] == "persons":
                person = e["person"]
                db.execute("INSERT INTO person_details VALUES(?,?,?,?,?,?)", (
                    entity_id, person["фамилия"], person["имя"], person["отчество"],
                    person["дата_рождения"], person["место_рождения"]))
                spans = e["source_spans"]
            for field, span in spans.items():
                if span is not None:
                    db.execute("INSERT INTO mentions(entity_id,original_text,start,end,note,field) VALUES(?,?,?,?,?,?)",
                        (entity_id, job["text"][span["start"]:span["end"]], span["start"], span["end"],
                         str(e.get("note", "")), field))
    return {"document_id": document_id, "already_saved": False}
